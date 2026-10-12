#!/usr/bin/env python3
"""Build download-copy DOCX from the isolated verb120 dataset."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt
root=Path(__file__).resolve().parents[3]
items=json.loads((root/'combined-expressions/verb120/data.json').read_text(encoding='utf-8'))
assert len(items)==120
out=root/'combined-expressions/verb120/TCF-120-verb-constructions.docx'
doc=Document()
doc.add_heading('TCF Canada · 动词句法结构 120 条',0)
doc.add_paragraph('口语 Tâche 2 / Tâche 3 · 核心 40 + 高频 50 + 进阶 30')
doc.add_paragraph('qch = quelque chose；qn = quelqu’un。Piper 在网页里朗读完整形式。')
for lvl,title in [('P1','核心 40'),('P2','高频 50'),('P3','进阶 30')]:
    doc.add_heading(f'{lvl} · {title}',level=1)
    for x in items:
        if x['level']!=lvl:continue
        p=doc.add_paragraph()
        p.add_run(f"{x['id']:03d} · {x['task']} · {x['pattern']}").bold=True
        p.add_run('\n'+x['zh']+'\n'+x['example'])
doc.save(out)
print('Wrote',out, 'rows',len(items))
