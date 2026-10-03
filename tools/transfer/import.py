from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import json,re,hashlib,collections
root=Path(__file__).resolve().parents[2]
import argparse
parser=argparse.ArgumentParser()
parser.add_argument('input_directory',type=Path)
inputs=parser.parse_args().input_directory
doc=Document(next(inputs.glob('*.docx')))
themes=['旅行与移民','职业与职场','教育与儿童','媒体与文化','环境与城市','科技与网络','饮食与健康','社会与公共事务','家庭与人际']
entries=[];notes={t:[] for t in themes};theme=None;level=None;table_index=0
verb=r'(?:^|[,，] )(?:se |s’|s\x27)?[a-zà-ÿ-]+(?:er|ir|re)\b'
for block in doc.element.body:
 if block.tag==qn('w:p'):
  text=''.join(block.itertext()) # use paragraph objects below for notes instead
 elif block.tag==qn('w:tbl'):
  table=doc.tables[table_index];theme=themes[table_index//6];table_index+=1
  for j,row in enumerate(table.rows[1:],1):
   lvl,fr,zh,questions,usage=[c.text for c in row.cells]
   parts=usage.split('；',1);tfr=parts[0];tzh=parts[1] if len(parts)>1 else ''
   assert lvl in ['A1','A2','B1','B2','C1','C2'] and all([fr,zh,questions,tfr,tzh]),(table_index,j)
   typ='动词搭配' if re.search(verb,fr) else '主题词汇'
   if re.search(r'\b(?:faire face|tenir compte|prendre en compte|avoir accès|trouver ses repères|mettre en|faire preuve)\b',fr):typ='固定搭配'
   ident='transfer-'+hashlib.sha256(f'{theme}|{lvl}|{fr}|{questions}|{usage}'.encode()).hexdigest()[:16]
   score=(3 if lvl in ['B1','B2'] else 2 if lvl=='C1' else 1 if lvl=='C2' else 0)+(2 if typ!='主题词汇' else 0)
   if re.search(r'expérience|compétence|concilier|formation|liens|logement|autonomie|qualité|public|ressources|information|confiance|partager|préserver|intégr|emploi|dialogue',fr):score+=3
   if re.search(r'biomimétisme|arômes|génériques|photocopieuse|bagages|embarquement|panda|SMS|extractive|spatiaux',fr,re.I):score-=4
   entries.append(dict(id=ident,theme=theme,level=lvl,fr=fr,zh=zh,questions=questions,t3=tfr,t3zh=tzh,usageOriginal=usage,type=typ,source='听力原文语料',t3Source='T3 迁移句 / 整理版',priority=score,table=table_index,row=j))
active=None
for p in doc.paragraphs:
 if p.text in themes: active=p.text
 elif p.text=='各等级原文件说明与复习优先级':active=None
 elif active and p.text:notes[active].append(p.text)
md=next(inputs.glob('*.md')).read_text();structures=[];section='';raw=[]
clean=lambda s:re.sub(r'\*|`','',s).strip()
for line in md.splitlines():
 if line.startswith('## '):section=line[3:]
 if not line.startswith('|') or re.match(r'^\|[-: ]+\|',line):continue
 cells=[s.strip() for s in line.strip('|').split('|')]
 if cells[0] in ['结构（规范写法）','搭配','功能','容易误用']:continue
 raw.append((section,cells))
 if section.startswith(('一、','二、')):
  fr,zh,source,example=map(clean,cells);typ='动词搭配';group='动词结构'
 elif section.startswith('三、'):
  function,fr,zh,example=map(clean,cells);source='整理版 · '+function;typ='连接结构';group='连接结构'
 elif section.startswith('四、'):
  wrong,fr,example=map(clean,cells);zh='易错提醒：'+wrong;source='Markdown · 六组易错结构';typ='固定搭配';group='易错结构'
 else:continue
 ident='transfer-structure-'+hashlib.sha256((group+fr).encode()).hexdigest()[:16]
 structures.append(dict(id=ident,fr=fr,zh=zh,questions=source,example=example,type=typ,group=group,source='T3 迁移句 / 整理版',level='跨等级',notes='规范化结构及原创例句，不是听力逐字摘录。'))
counts=collections.Counter(e['theme'] for e in entries);levelcounts=collections.Counter(e['level'] for e in entries)
exact=collections.Counter((e['theme'],e['level'],e['fr'],e['zh'],e['questions'],e['t3']) for e in entries)
sim=collections.defaultdict(list)
for e in entries:sim[e['fr'].casefold().replace('’',"'")].append(e['id'])
audit=dict(total=len(entries),themes=dict(counts),levels=dict(levelcounts),structures=len(structures),structureGroups=dict(collections.Counter(s['group'] for s in structures)),exactDuplicates=sum(n-1 for n in exact.values()),repeatedExpressions={k:v for k,v in sim.items() if len(v)>1},policy='主题、等级、题号、中文、T3用法不同则保留独立语境；辅助库为规范化结构，不与听力词条混为一条。',sourceFiles=[p.name for p in inputs.iterdir()])
data=dict(themes=themes,entries=entries,structures=structures,themeNotes=notes,documentNotes=[p.text for p in doc.paragraphs if p.text],markdownOriginal=md,audit=audit)
(root/'combined-expressions/transfer-data.js').write_text('window.TRANSFER_RESTORE=(()=>{try{return JSON.parse(localStorage.getItem("TCF-COMBINED-EXPRESSIONS-v1"))?.module===11}catch{return false}})();\nwindow.TRANSFER_DATA='+json.dumps(data,ensure_ascii=False,separators=(',',':'))+';\n')
(root/'combined-expressions/transfer-audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
print(json.dumps(audit,ensure_ascii=False,indent=2))
