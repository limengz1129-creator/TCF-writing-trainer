#!/usr/bin/env python3
"""Export nine C1 topics to nine Word files: hierarchy + bilingual argument paragraphs only."""
import json, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path(__file__).resolve().parents[1]/'eo-tache3-corpus'
OUT=Path(__file__).resolve().parents[1]/'exports'/'t3-c1-paragraphs'
OUT.mkdir(parents=True, exist_ok=True)
TOPICS=[
('travel','旅行与移民'),('career','职业与职场'),('education','教育与儿童'),
('environment','环境与城市'),('health','饮食与健康'),('media','文化与媒体'),
('tech','科技与网络'),('social','社会与公共事务'),('family','家庭与人际')]
ZH=re.compile(r'[\u3400-\u9fff]')
FR=re.compile(r'[A-Za-zÀ-ÿ]')
def parse(path):
    src=path.read_text(encoding='utf-8')
    return json.loads(src[src.index('=')+1:].rstrip().rstrip(';'))
def bilingual(text):
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    out=[]; french=[]
    started=False
    for line in lines:
        if FR.search(line) and not ZH.search(line):
            french.append(line)
            started=True
        elif ZH.search(line) and started:
            if french:
                out.append(('fr',' '.join(french)))
                french=[]
            out.append(('zh',line))
    if french:out.append(('fr',' '.join(french)))
    return out
for slug,title in TOPICS:
    entries=parse(ROOT/f'c1-{slug}-data.js')
    doc=Document()
    section=doc.sections[0]
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    section.left_margin=Cm(2.25);section.right_margin=Cm(2.25)
    styles=doc.styles
    normal=styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(4)
    for s,size in [('Title',20),('Heading 1',15),('Heading 2',13),('Heading 3',11)]:
        styles[s].font.name='Arial';styles[s].font.size=Pt(size)
        styles[s].font.color.rgb=RGBColor(29,55,82)
    frstyle=styles.add_style('French Paragraph',WD_STYLE_TYPE.PARAGRAPH)
    frstyle.base_style=normal;frstyle.font.name='Arial';frstyle.font.size=Pt(10.5)
    frstyle.paragraph_format.space_before=Pt(7);frstyle.paragraph_format.space_after=Pt(3)
    zhstyle=styles.add_style('Chinese Translation',WD_STYLE_TYPE.PARAGRAPH)
    zhstyle.base_style=normal;zhstyle.font.name='Microsoft YaHei';zhstyle.font.size=Pt(10)
    zhstyle.font.color.rgb=RGBColor(75,80,87);zhstyle.paragraph_format.space_after=Pt(4)
    doc.add_heading(title,0)
    group=None;mother=None
    for e in entries:
        g=e.get('group') or '母论点'
        m=e.get('mother') or ''
        if g!=group:
            group=g;mother=None
            doc.add_heading(g,level=1)
        if m!=mother:
            mother=m;doc.add_heading(m,level=2)
        label=f"{e['id']} · {e.get('scene','情境')}"
        doc.add_heading(label,level=3)
        for lang,line in bilingual(e.get('text','')):
            doc.add_paragraph(line,style='French Paragraph' if lang=='fr' else 'Chinese Translation')
    dest=OUT/f'{title}｜C1完整论段.docx'
    doc.save(dest)
    print(f'{title}: {len(entries)} paragraphs -> {dest}')
